# -*- coding: utf-8 -*-
"""Gera os dados de exemplo: tudo aqui e ficticio, inventado por este script.

    python dados-exemplo/gerar.py

Empresas, CNPJs, sites, trechos e historico de vendas sao inventados. Os CNPJs
comecam por 99 e tem digito verificador errado de proposito, para nunca
coincidir com uma empresa real. Os dominios terminam em `.exemplo`, que nao
existe na internet.

Saidas, nesta pasta:
  base-exemplo.xlsx      prospects e historico de oportunidades (calibra o segmento)
  lista.jsonl            a lista de origem: nome, cidade, site e cargo de contato
  casamento.json         cada empresa da lista com o CNPJ (ficticio)
  receita-exemplo.json   o que a Receita diria de cada CNPJ (ficticio)
  cnaes-exemplo.json     descricao dos CNAEs usados
  extracao.jsonl         o que o modelo teria lido no site, com o trecho
"""
from __future__ import annotations

import json
import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl

AQUI = Path(__file__).resolve().parent
HOJE = date(2026, 10, 17)
rnd = random.Random(2026)

NOMES = ["Aurora", "Boreal", "Cobalto", "Duna", "Estuario", "Granito", "Horizonte", "Iris",
         "Jacaranda", "Kepler", "Lotus", "Moncao", "Nimbus", "Oasis", "Pampa", "Quartzo",
         "Restinga", "Safira", "Tamarindo", "Ubiratan", "Vereda", "Xisto", "Zenite", "Alvorada",
         "Bromelia", "Cerrado", "Delta", "Esmeralda", "Figueira", "Gaivota", "Hibisco", "Ipe",
         "Jatoba", "Lince", "Mandacaru", "Nascente", "Orquidea", "Paineira", "Quaresmeira", "Rubi"]
SEGMENTOS = {  # segmento -> (CNAE principal, descricao, ramo no nome)
    "Indústria": ("2610800", "Fabricação de componentes eletrônicos", "Componentes"),
    "Tecnologia": ("6201501", "Desenvolvimento de programas de computador sob encomenda", "Sistemas"),
    "Logística": ("4930202", "Transporte rodoviário de carga, intermunicipal e interestadual", "Transportes"),
    "Saúde": ("8610101", "Atividades de atendimento hospitalar", "Saúde"),
    "Varejo": ("4753900", "Comércio varejista de eletrodomésticos e equipamentos de áudio e vídeo", "Magazine"),
    "Serviços financeiros": ("6422100", "Bancos múltiplos, com carteira comercial", "Crédito"),
    "Energia": ("3511501", "Geração de energia elétrica", "Energia"),
    "Educação": ("8532500", "Educação superior, graduação e pós-graduação", "Educação"),
}
CIDADES = [("São Paulo", "SP"), ("Campinas", "SP"), ("Curitiba", "PR"), ("Porto Alegre", "RS"),
           ("Belo Horizonte", "MG"), ("Recife", "PE"), ("Salvador", "BA"), ("Goiânia", "GO"),
           ("Florianópolis", "SC"), ("Fortaleza", "CE")]
UFS = ["SP", "RJ", "MG", "PR", "RS", "SC", "BA", "PE", "GO", "CE", "DF", "ES", "AM", "PA"]
NIVEIS = ["Alto", "Médio", "Baixo", "Não identificado"]
RENOV = ["0-3 meses", "4-6 meses", "7-12 meses", ">12 meses", "Não identificado"]
CARGOS = ["Gerente de Recursos Humanos", "Coordenadora de Gente e Gestão", "Diretor de Pessoas",
          "Analista de Benefícios", "Gerente de Marketing", "Comprador"]


def empresas(n: int = 40) -> list[dict]:
    saida = []
    for i in range(n):
        seg = list(SEGMENTOS)[i % len(SEGMENTOS)]
        nome = NOMES[i]
        colab = int(min(9000, max(40, rnd.lognormvariate(6.3, 1.0))))
        cidade, uf = rnd.choice(CIDADES)
        multi = colab > 600 and rnd.random() < 0.7
        saida.append({
            "id": f"EX-{i + 1:04d}", "codigo": str(i + 1), "nome": nome, "segmento": seg,
            "razao": f"{nome.upper()} {SEGMENTOS[seg][2].upper()} EXEMPLO LTDA",
            "fantasia": f"{nome} {SEGMENTOS[seg][2]}",
            "dominio": f"{nome.lower()}.exemplo", "colab": colab, "cidade": cidade, "uf": uf,
            "rh": colab > 250 and rnd.random() < 0.8, "multi": multi,
            "cresc": rnd.choice(NIVEIS), "eb": rnd.choice(NIVEIS), "dor": rnd.choice(NIVEIS),
            "renov": rnd.choice(RENOV), "ativa": rnd.random() > 0.05,
            "cnpj_basico": f"99{i + 1:06d}",
        })
    return saida


def base_xlsx(emps: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leia-me"
    ws.append(["Base de exemplo, inteiramente ficticia, gerada por dados-exemplo/gerar.py."])
    p = wb.create_sheet("Prospects")
    p.append(["Prospects (ficticio)"])
    p.append([])
    p.append(["ID Empresa", "Empresa", "Domínio fictício", "Segmento", "Colaboradores", "UF",
              "RH estruturado", "Multi-região", "Sinal crescimento", "Employer branding",
              "Dor retenção", "Renovação provável", "Fonte permitida", "Status"])
    for e in emps:
        p.append([e["id"], e["fantasia"] + " (fictícia)", e["dominio"], e["segmento"], e["colab"], e["uf"],
                  "Sim" if e["rh"] else "Não", "Sim" if e["multi"] else "Não", e["cresc"], e["eb"],
                  e["dor"], e["renov"], "Site institucional", "Ativa" if e["ativa"] else "Inativa"])
    o = wb.create_sheet("Oportunidades_Historico")
    o.append(["Oportunidades (ficticio)"])
    o.append([])
    o.append(["ID Oportunidade", "ID Empresa", "Empresa", "Resultado", "Ciclo (dias)", "Ticket estimado",
              "Motivo/Status", "Sponsor/Decisor", "Origem"])
    for k in range(48):
        e = rnd.choice(emps)
        fit = (e["rh"] + e["multi"] + (e["eb"] == "Alto") + (200 <= e["colab"] <= 5000)) / 4
        r = rnd.random()
        res = "Em andamento" if r < 0.15 else ("Ganha" if rnd.random() < 0.25 + 0.5 * fit else "Perdida")
        motivo = {"Ganha": "Fechou contrato", "Perdida": rnd.choice(["Preço", "Sem prioridade", "Concorrente"]),
                  "Em andamento": "Proposta enviada"}[res]
        o.append([f"OP-{k + 1:03d}", e["id"], e["fantasia"] + " (fictícia)", res, rnd.randint(20, 160),
                  rnd.choice([60000, 85000, 110000, 140000]), motivo, rnd.choice(CARGOS[:4]),
                  rnd.choice(["Inbound", "Outbound", "Indicação"])])
    wb.save(AQUI / "base-exemplo.xlsx")


def _campo(valor, url, trecho, escopo=None):
    c = {"valor": valor, "url": url, "recortado": False, "tipo": "extraido",
         "fonte": "site da empresa (exemplo ficticio)", "trecho": trecho, "trecho_mascarado": False}
    if escopo:
        c["escopo"] = escopo
    return c


def extracao(e: dict) -> dict:
    site = f"https://www.{e['dominio']}"
    campos = {}
    if e["rh"]:
        campos["rh_estruturado"] = _campo(True, f"{site}/trabalhe-conosco",
            f"Nossa área de Gente e Gestão cuida de recrutamento, treinamento e benefícios dos {e['colab']} colaboradores.")
    if e["multi"]:
        campos["multi_localidade"] = _campo(True, f"{site}/sobre",
            f"Estamos presentes em {rnd.randint(3, 9)} estados, com unidades e equipes regionais.")
    if rnd.random() < 0.6:
        campos["colaboradores"] = _campo(e["colab"], f"{site}/sobre",
            f"Somos mais de {e['colab']:,} pessoas trabalhando juntas.".replace(",", "."))
    if e["cresc"] == "Alto":
        mes = (HOJE - timedelta(days=rnd.randint(20, 200)))
        campos["sinal_crescimento"] = _campo(True, f"{site}/noticias",
            f"Em {mes.strftime('%d/%m/%Y')} inauguramos uma nova unidade e abrimos vagas em várias áreas.")
    if e["eb"] == "Alto":
        campos["employer_branding"] = _campo(True, f"{site}/carreiras",
            "Fomos reconhecidos como um dos melhores lugares para trabalhar da região em 2026.")
    elif e["eb"] == "Médio":
        campos["employer_branding"] = _campo(True, f"{site}/carreiras",
            "Conheça as histórias de quem trabalha aqui e o que torna nossa cultura especial.")
    if e["dor"] == "Alto":
        campos["dor_retencao"] = _campo(True, f"{site}/blog/retencao",
            "Nosso desafio é reduzir o turnover e reter os talentos que formamos.")
    if rnd.random() < 0.08:
        return {"codigo": e["codigo"], "razao_social": e["razao"], "dominio": e["dominio"],
                "paginas": [{"url": site + "/", "resultado": "RobotsIndisponivel"}],
                "erro": "nenhuma pagina legivel"}
    return {"codigo": e["codigo"], "razao_social": e["razao"], "dominio": e["dominio"],
            "paginas": [{"url": site + "/", "resultado": "ok", "caracteres": rnd.randint(2000, 12000)}],
            "campos": campos, "descartados": [], "truncado": False, "erro": None}


def receita(e: dict) -> dict:
    n = max(1, min(60, e["colab"] // 150 + (rnd.randint(2, 12) if e["multi"] else 0)))
    ufs = sorted({e["uf"], *(rnd.sample(UFS, rnd.randint(2, 6)) if e["multi"] else [])})
    abertos = rnd.randint(1, 6) if e["cresc"] == "Alto" else (rnd.randint(0, 2) if e["cresc"] == "Médio" else 0)
    ultima = None
    if abertos:
        cid, uf = rnd.choice(CIDADES)
        ultima = {"data_inicio_atividade": (HOJE - timedelta(days=rnd.randint(30, 600))).isoformat(),
                  "uf": uf, "municipio": cid.upper()}
    return {"estabelecimentos": n, "ufs": ufs, "uf_matriz": e["uf"],
            "cnae_principal": SEGMENTOS[e["segmento"]][0], "abertos_24m": abertos,
            "filial_mais_recente": ultima, "razao_social": e["razao"],
            "porte_receita": "05" if e["colab"] > 100 else "03"}


def main():
    emps = empresas()
    base_xlsx(emps)
    lista = emps[:24]
    fonte = "lista de origem (exemplo ficticio)"
    with open(AQUI / "lista.jsonl", "w", encoding="utf-8") as f:
        for e in lista:
            url = f"https://lista.exemplo/empresa/{e['codigo']}"
            campos = {"razao_social": e["razao"], "contato_cargo": rnd.choice(CARGOS),
                      "municipio": e["cidade"], "uf": e["uf"], "site": f"https://www.{e['dominio']}",
                      "dominio": e["dominio"]}
            prov = {k: {"fonte": fonte, "url": url, "tipo": "observado"} for k in campos}
            f.write(json.dumps({"codigo": e["codigo"], "campos": campos, "proveniencia": prov},
                               ensure_ascii=False) + "\n")
    cas = []
    for e in lista:
        cnpj = e["cnpj_basico"] + "000100"   # digito verificador errado de proposito
        cas.append({"codigo": e["codigo"], "lista": e["razao"], "cidade": e["cidade"], "uf": e["uf"],
                    "status": "confiavel", "cnpj": cnpj,
                    "candidato": {"cnpj": cnpj, "razao_social": e["razao"],
                                  "municipio": e["cidade"].upper(), "uf": e["uf"]}})
    (AQUI / "casamento.json").write_text(json.dumps(cas, ensure_ascii=False, indent=1), encoding="utf-8")
    (AQUI / "receita-exemplo.json").write_text(
        json.dumps({e["cnpj_basico"]: receita(e) for e in lista}, ensure_ascii=False, indent=1), encoding="utf-8")
    (AQUI / "cnaes-exemplo.json").write_text(
        json.dumps({c: d for c, d, _ in SEGMENTOS.values()}, ensure_ascii=False, indent=1), encoding="utf-8")
    with open(AQUI / "extracao.jsonl", "w", encoding="utf-8") as f:
        for e in lista:
            f.write(json.dumps(extracao(e), ensure_ascii=False) + "\n")
    print(f"{len(emps)} prospects, {len(lista)} empresas na lista de exemplo")


if __name__ == "__main__":
    main()
